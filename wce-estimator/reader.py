#!/usr/bin/env python3
"""
Author: Viquar Khan (ORCID 0009-0008-3592-4162).

Multi-tab Excel reader for the WCE Estimator.

Reads every worksheet ("tab") of an uploaded workbook. Each tab is treated as a
subsystem / module / project area; each data row is a work item.

Recognized columns (case-insensitive, flexible aliases). Only "work item" and an
effort column are required; everything else is inferred from the paper's priors:

    Work Item / Task / Feature / Component      (required) - the label
    Work Class / Category / Type                (optional) - overrides classification
    Baseline / Effort / Estimate / PM / Days    (required) - baseline effort number
    AI Difficulty / Difficulty / Tier           (optional) - Low / Medium / High
    Essential / Non-accelerable                 (optional) - Yes/No/True/False
    a_low / a_mode / a_high  (or) Multiplier    (optional) - explicit override band

An optional sheet named "Config" / "Settings" / "Meta" may carry key/value rows
such as:  Team Size = 2 ,  Unit = person-months ,  Claim Months = 1.

Resolution priority for the multiplier band:
    1. explicit a_low/a_mode/a_high (or a single Multiplier value)
    2. AI-difficulty tier band
    3. work-class default band
"""

from __future__ import annotations

import math
import re
import openpyxl

from .priors import (
    normalize_difficulty, classify_work, resolve_band, is_essential_label,
    DEFAULT_CLASS,
)
from .engine import WorkItem

# ---- column alias maps -----------------------------------------------------
ALIASES = {
    "name": ["work item", "work item name", "workitem", "item", "line item",
             "task", "task name", "feature", "feature name", "component",
             "module", "activity", "work", "work description", "description",
             "deliverable", "epic", "epic / story", "epic/story",
             "epic / stories", "story", "user story", "epic name", "story name",
             "requirement", "title", "summary"],
    "work_class": ["work class", "workclass", "class", "category", "type", "kind",
                   "discipline"],
    "baseline": ["baseline", "baseline effort", "effort", "effort (days)",
                 "effort (hrs)", "effort (hours)", "effort estimate", "estimate",
                 "estimate (days)", "base effort", "person-months", "person months",
                 "pm", "person-days", "person days", "days", "man-days", "man days",
                 "mandays", "hours", "hrs", "story points", "points", "pts", "size",
                 "baseline pm", "baseline (pm)", "baseline months",
                 "baseline effort (pm)", "total", "total days", "total effort",
                 "dev days", "dev-days", "developer days", "developer-days",
                 "loe", "effort days", "duration"],
    "difficulty": ["ai difficulty", "difficulty", "ai-difficulty", "tier",
                   "ai tier", "acceleration", "ai resistance"],
    "essential": ["essential", "non-accelerable", "nonaccelerable", "essence",
                  "must-human", "human-only"],
    "a_low": ["a_low", "a low", "alow", "multiplier low", "mult low", "low"],
    "a_mode": ["a_mode", "a mode", "amode", "multiplier", "mult", "a", "mode",
               "effort multiplier"],
    "a_high": ["a_high", "a high", "ahigh", "multiplier high", "mult high", "high"],
}

CONFIG_SHEETS = {"config", "settings", "meta", "parameters", "params"}
# Documentation / cover sheets that carry no work items; skipped silently.
IGNORE_SHEETS = {"instructions", "readme", "read me", "help", "cover", "about",
                 "notes", "legend", "guide", "how to", "howto"}
TRUE_WORDS = {"yes", "y", "true", "1", "essential", "t"}
# NOTE: an empty/blank cell is intentionally NOT here -> _to_bool returns None so
# the work-class default decides essential (blank must not force non-essential).
FALSE_WORDS = {"no", "n", "false", "0", "f"}
# Rows whose label is a rollup / total, not a work item; excluded to avoid
# double-counting the workbook's own subtotals.
TOTAL_LABELS = {"total", "totals", "subtotal", "sub-total", "grand total", "sum",
                "total:", "grand total:", "sum:"}


def _norm(s) -> str:
    return str(s).strip().lower() if s is not None else ""


# Phase-2 fallback: word-boundary tokens for headers that are not an EXACT alias
# (e.g. "Task Name", "Effort (days)", "LOE in days"). Only reasonably specific
# tokens (>= 3 chars) are listed to avoid false hits. Baseline is resolved before
# name so "Story Points" is read as effort, not as the item name.
CONTAINS = {
    "baseline": ["baseline", "effort", "estimate", "person-month", "person month",
                 "person-day", "person day", "dev day", "developer day",
                 "story point", "points", "hours", "days", "man-day", "man day",
                 "manday", "size", "loe", "duration", "total"],
    "name": ["work item", "work-item", "workitem", "task", "feature", "story",
             "epic", "deliverable", "activity", "component", "module",
             "requirement", "description", "item", "title", "summary"],
    "work_class": ["work class", "category", "discipline", "kind"],
    "difficulty": ["difficulty", "ai tier", "ai resistance", "acceleration"],
    "essential": ["essential", "non-accelerable", "essence"],
}
_CONTAINS_PATTERNS = {
    field: [(re.compile(r"\b" + re.escape(tok) + r"\b"), tok) for tok in toks]
    for field, toks in CONTAINS.items()
}


def _match_header(header_cells):
    """
    Return {canonical_field: column_index} for a header row.

    Phase 1: exact alias match (precise). Phase 2: for still-missing name/baseline
    and other fields, a word-boundary token match so common variants like
    "Task Name" or "Effort (days)" are recognized without a rigid template.
    """
    norm = [_norm(c) for c in header_cells]
    mapping = {}
    used = set()

    # Phase 1: exact alias equality
    for idx, h in enumerate(norm):
        if not h or idx in used:
            continue
        for field, names in ALIASES.items():
            if field in mapping:
                continue
            if h in names:
                mapping[field] = idx
                used.add(idx)
                break

    # Phase 2: word-boundary token fallback (baseline first to avoid collisions)
    for field in ("baseline", "name", "work_class", "difficulty", "essential"):
        if field in mapping:
            continue
        for idx, h in enumerate(norm):
            if not h or idx in used:
                continue
            if any(pat.search(h) for pat, _tok in _CONTAINS_PATTERNS[field]):
                mapping[field] = idx
                used.add(idx)
                break
    return mapping


def _to_float(v):
    """Parse a number from a cell, tolerant of thousands separators, European
    decimal commas, and trailing units (e.g. '1,234.5', '1,5', '5 days')."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if not (isinstance(v, float) and math.isnan(v)) else None
    s = str(v).strip()
    if not s:
        return None
    # normalize commas: thousands separator vs decimal comma
    if "," in s and "." in s:
        s = s.replace(",", "")                          # 1,234.5 -> 1234.5
    elif "," in s:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", s):        # 1,000 / 12,345,678 -> thousands
            s = s.replace(",", "")
        elif re.fullmatch(r"\d+,\d{1,2}", s):           # 1,5 -> decimal comma
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        m = re.search(r"-?\d+(?:\.\d+)?", s)             # '5 days', '~2', '3 pts'
        return float(m.group()) if m else None


def _to_bool(v):
    w = _norm(v)
    if w in TRUE_WORDS:
        return True
    if w in FALSE_WORDS:
        return False
    return None


def _read_config(ws):
    cfg = {}
    for row in ws.iter_rows(values_only=True):
        if row is None or len(row) < 2:
            continue
        key, val = _norm(row[0]), row[1]
        if not key:
            continue
        if "team" in key:
            f = _to_float(val)
            if f:
                cfg["team_size"] = f
        elif key in ("unit", "effort unit"):
            cfg["unit"] = str(val).strip()
        elif "claim" in key:
            f = _to_float(val)
            if f:
                cfg["claim_calendar"] = f
        elif "working day" in key or "days per month" in key or "day/month" in key:
            f = _to_float(val)
            if f:
                cfg["days_per_month"] = f
        elif "project" in key or "name" == key:
            cfg["project"] = str(val).strip()
    return cfg


def read_workbook(path, include_sheets=None, default_difficulty=None):
    """
    Parse a workbook into (items, config, warnings).

    include_sheets : optional list of sheet names to restrict data parsing to
                     (case-insensitive). Useful when a workbook holds several
                     VIEWS of the same scope (e.g. an Epic summary and a Story
                     detail) that must not be summed together. The Config sheet
                     is always read regardless.

    default_difficulty : optional 'low'/'medium'/'high' applied to rows that
                     could NOT be classified (no work-class match, no difficulty
                     column, no explicit multiplier). Use 'medium' for coarse
                     epic-level lists where each row bundles code + test +
                     integration + review, so the aggressive pure-coding default
                     would overstate AI benefit.

    items    : list[WorkItem]
    config   : dict (team_size, unit, claim_calendar, project)
    warnings : list[str]
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    items: list[WorkItem] = []
    warnings: list[str] = []
    config = {"team_size": 1.0, "unit": "person-months",
              "claim_calendar": None, "project": None, "days_per_month": 21.0}
    want = {s.strip().lower() for s in include_sheets} if include_sheets else None
    default_tier = normalize_difficulty(default_difficulty) if default_difficulty else None

    for ws in wb.worksheets:
        sheet_name = ws.title
        if _norm(sheet_name) in CONFIG_SHEETS:
            config.update(_read_config(ws))
            continue
        if _norm(sheet_name) in IGNORE_SHEETS:
            continue
        if want is not None and _norm(sheet_name) not in want:
            continue

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # find the header row (first row that maps a name + baseline column)
        header_idx, mapping = None, None
        for i, row in enumerate(rows[:25]):
            m = _match_header(row)
            if "name" in m and "baseline" in m:
                header_idx, mapping = i, m
                break
        if mapping is None:
            warnings.append(f"[{sheet_name}] skipped: no 'Work Item' + effort "
                            f"columns found.")
            continue

        sheet_items = 0
        for row in rows[header_idx + 1:]:
            if row is None:
                continue
            name = row[mapping["name"]] if mapping["name"] < len(row) else None
            name = str(name).strip() if name is not None else ""
            baseline = _to_float(row[mapping["baseline"]]) if mapping["baseline"] < len(row) else None
            if not name or baseline is None or baseline <= 0:
                continue
            nl = name.strip().lower()
            if nl in TOTAL_LABELS or nl.startswith("total ") or nl.startswith("subtotal"):
                continue

            def cell(field):
                if field in mapping and mapping[field] < len(row):
                    return row[mapping[field]]
                return None

            work_class_label = cell("work_class")
            difficulty_raw = cell("difficulty")
            tier = normalize_difficulty(difficulty_raw)

            # classification uses explicit work-class label if present, else the item name
            class_key = classify_work(work_class_label or name)

            a_low = _to_float(cell("a_low"))
            a_mode = _to_float(cell("a_mode"))
            a_high = _to_float(cell("a_high"))

            # essential resolution
            ess_cell = _to_bool(cell("essential"))

            source = ""
            eff_tier = tier
            if a_mode is not None:
                # explicit multiplier(s)
                lo = a_low if a_low is not None else a_mode * 0.85
                hi = a_high if a_high is not None else a_mode * 1.20
                lo, hi = min(lo, a_mode), max(hi, a_mode)
                base_essential = is_essential_label(work_class_label or name)
                source = "explicit multiplier"
            else:
                # unclassified + no difficulty given -> fall back to the caller's
                # default assumption (keeps coarse epic lists from over-crediting AI)
                if eff_tier is None and default_tier and class_key == DEFAULT_CLASS \
                        and not (work_class_label and str(work_class_label).strip()):
                    eff_tier = default_tier
                lo, md, hi, base_essential = resolve_band(class_key, eff_tier)
                a_mode = md
                a_low, a_high = lo, hi
                source = (f"difficulty:{eff_tier}" if eff_tier else f"class:{class_key}")

            essential = ess_cell if ess_cell is not None else base_essential

            items.append(WorkItem(
                name=name,
                subsystem=sheet_name,
                baseline=baseline,
                a_low=float(a_low if a_low is not None else lo),
                a_mode=float(a_mode),
                a_high=float(a_high if a_high is not None else hi),
                essential=bool(essential),
                work_class=class_key,
                difficulty=(eff_tier or ""),
                source=source,
            ))
            sheet_items += 1

        if sheet_items == 0:
            warnings.append(f"[{sheet_name}] header found but no valid data rows.")

    wb.close()
    if not items:
        raise ValueError("No valid work items found in any sheet. Each data sheet "
                         "needs a 'Work Item' column and an effort/baseline column.")
    return items, config, warnings
