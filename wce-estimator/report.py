#!/usr/bin/env python3
"""
Author: Viquar Khan (ORCID 0009-0008-3592-4162).

Report builders for the WCE Estimator.

    build_excel_report(res, path)  -> multi-tab .xlsx workbook
    build_html_report(res)         -> self-contained HTML string (embedded charts)
"""

from __future__ import annotations

import html

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import figures

HEAD_FILL = PatternFill("solid", fgColor="1e3a8a")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1e3a8a")
SUB_FONT = Font(bold=True, size=11, color="1e3a8a")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cal_unit(unit):
    return figures._cal_unit(unit)


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header_row(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = BORDER


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_excel_report(res, path):
    wb = openpyxl.Workbook()
    cal = _cal_unit(res.unit)

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "WCE Estimator - Final Estimation Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Work-Class Estimation with the Khan Ceiling (Viquar Khan)"
    ws["A2"].font = Font(italic=True, color="6b7280")

    rows = [
        ("Metric", "Value"),
        (f"Baseline effort ({res.unit})", res.baseline_effort),
        (f"AI-adjusted effort ({res.unit})", res.adjusted_effort),
        ("Adjusted effort ratio r", res.r),
        ("Whole-project speedup S", f"{res.speedup:.2f}x"),
        ("Net effort reduction", f"{res.net_reduction_pct:.1f}%"),
        ("Team size", res.team_size),
        (f"Baseline calendar ({cal})", res.baseline_calendar),
        (f"AI-adjusted calendar ({cal})", res.adjusted_calendar),
        (f"Calendar time saved ({cal})", res.calendar_saved),
        ("", ""),
        ("Essential fraction phi", res.phi),
        ("Khan Ceiling (strict)", _fmt_ceiling(res.ceiling_strict)),
        ("Khan Ceiling (generous)", _fmt_ceiling(res.ceiling_generous)),
        (f"Absolute floor ({cal}) - nothing beats this", res.floor_calendar),
        ("", ""),
        (f"Conditional P10-P90 calendar ({cal})", f"{res.cond_p10_cal:.2f} - {res.cond_p90_cal:.2f}"),
        (f"Total (+scope) P10-P90 calendar ({cal})", f"{res.total_p10_cal:.2f} - {res.total_p90_cal:.2f}"),
        ("", ""),
        ("Flat-discount misallocation (P2)", f"{res.misallocated_effort:.2f} {res.unit}"),
        ("Codegen Illusion gap (P3)", f"{res.codegen_gap:.2f}x"),
        ("  most-accelerated class", res.focus_class),
    ]
    start = 4
    for i, (k, v) in enumerate(rows):
        r = start + i
        kc = ws.cell(row=r, column=1, value=k)
        vc = ws.cell(row=r, column=2, value=v)
        if i == 0:
            kc.fill = HEAD_FILL; kc.font = HEAD_FONT
            vc.fill = HEAD_FILL; vc.font = HEAD_FONT
        elif k and not v == "":
            kc.font = Font(bold=("Khan" in k or "speedup" in k.lower() or "adjusted calendar" in k.lower()))
    if res.claim:
        r = start + len(rows) + 1
        ws.cell(row=r, column=1, value="Claim check").font = SUB_FONT
        claim_rows = [
            (f"Claimed calendar ({cal})", res.claim["claim_calendar"]),
            ("Implied speedup", f"{res.claim['implied_speedup']:.2f}x"),
            ("Exceeds generous ceiling?", res.claim["exceeds_generous_ceiling"]),
            ("P(<= claim) total", res.claim["prob_at_or_below_total"]),
            ("Verdict", res.claim["verdict"]),
        ]
        for j, (k, v) in enumerate(claim_rows):
            ws.cell(row=r + 1 + j, column=1, value=k)
            vc = ws.cell(row=r + 1 + j, column=2, value=str(v))
            if k == "Verdict":
                vc.font = Font(bold=True,
                               color="dc2626" if "REJECT" in str(v) or "IMPLAUSIBLE" in str(v) else "16a34a")
    _autofit(ws, [46, 26])

    # ---- Work Items sheet ----
    ws = wb.create_sheet("Work Items")
    headers = ["Subsystem", "Work Item", "Work Class", "AI Difficulty", "Essential",
               f"Baseline ({res.unit})", "a_low", "a_mode", "a_high",
               f"Adjusted ({res.unit})", "Basis"]
    _header_row(ws, 1, headers)
    for i, it in enumerate(res.items, start=2):
        vals = [it.subsystem, it.name, it.work_class, it.difficulty or "-",
                "Yes" if it.essential else "No", round(it.baseline, 3),
                round(it.a_low, 3), round(it.a_mode, 3), round(it.a_high, 3),
                round(it.baseline * it.a_mode, 3), it.source]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if it.essential and j == 5:
                c.font = Font(bold=True, color="dc2626")
    # totals
    tr = len(res.items) + 2
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=tr, column=6, value=round(res.baseline_effort, 3)).font = Font(bold=True)
    ws.cell(row=tr, column=8, value=round(res.r, 3)).font = Font(bold=True)
    ws.cell(row=tr, column=10, value=round(res.adjusted_effort, 3)).font = Font(bold=True)
    _autofit(ws, [22, 40, 16, 12, 10, 16, 8, 9, 8, 16, 20])
    ws.freeze_panes = "A2"

    # ---- Subsystems sheet ----
    ws = wb.create_sheet("Subsystems")
    headers = ["Subsystem", "Items", f"Baseline ({res.unit})", f"Adjusted ({res.unit})",
               "Speedup", "Net reduction %", "Essential %",
               f"Baseline cal ({cal})", f"Adjusted cal ({cal})"]
    _header_row(ws, 1, headers)
    for i, s in enumerate(res.subsystems, start=2):
        vals = [s["subsystem"], s["items"], s["baseline"], s["adjusted"],
                f"{s['speedup']:.2f}x", s["net_reduction_pct"], s["essential_pct"],
                s["baseline_calendar"], s["adjusted_calendar"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v).border = BORDER
    _autofit(ws, [26, 7, 16, 16, 10, 15, 12, 16, 16])
    ws.freeze_panes = "A2"

    # ---- Flat-Discount Error (P2) sheet ----
    ws = wb.create_sheet("Flat-Discount Error")
    ws["A1"] = ("P2: a uniform discount misallocates effort across classes even when "
                "its total is right.")
    ws["A1"].font = SUB_FONT
    ws[f"A2"] = f"Total misallocation M = {res.misallocated_effort:.2f} {res.unit}"
    ws["A2"].font = Font(bold=True, color="dc2626")
    headers = ["Work Item", "Share f_i", "a_mode",
               f"Flat budget ({res.unit})", f"True need ({res.unit})",
               f"Misallocation ({res.unit})"]
    _header_row(ws, 4, headers)
    for i, m in enumerate(res.misallocation_rows, start=5):
        vals = [m["item"], m["share"], m["a_mode"], m["flat_effort"],
                m["true_effort"], m["misallocation"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j == 6:
                c.font = Font(color="dc2626" if v > 0 else "16a34a")
    _autofit(ws, [40, 12, 10, 18, 18, 20])

    # ---- Difficulty Mix sheet ----
    ws = wb.create_sheet("Difficulty Mix")
    _header_row(ws, 1, ["AI Difficulty", "Items", f"Baseline ({res.unit})", "Baseline %"])
    for i, (k, v) in enumerate(sorted(res.difficulty_mix.items()), start=2):
        for j, val in enumerate([k, v["count"], v["baseline"], v["baseline_pct"]], start=1):
            ws.cell(row=i, column=j, value=val).border = BORDER
    _autofit(ws, [18, 8, 18, 12])

    if res.warnings:
        ws = wb.create_sheet("Notes")
        ws["A1"] = "Parser notes / warnings"
        ws["A1"].font = SUB_FONT
        for i, w in enumerate(res.warnings, start=2):
            ws.cell(row=i, column=1, value=w)
        _autofit(ws, [100])

    wb.save(path)
    return path


def _fmt_ceiling(v):
    return "unbounded" if v == float("inf") else round(v, 2)


# ---------------------------------------------------------------------------
# HTML (self-contained)
# ---------------------------------------------------------------------------
_STYLE = """<style>
  :root { --blue:#1e3a8a; --ink:#0f172a; --muted:#64748b; --line:#e2e8f0; }
  * { box-sizing:border-box; }
  body { font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;
         color:var(--ink); margin:0; background:#f8fafc; line-height:1.5; }
  header { background:var(--blue); color:#fff; padding:28px 32px; }
  header h1 { margin:0 0 4px; font-size:22px; }
  header p { margin:0; opacity:.85; font-size:13px; }
  main { max-width:1040px; margin:0 auto; padding:24px 32px 64px; }
  .kpis { display:flex; flex-wrap:wrap; gap:14px; margin:20px 0 8px; }
  .kpi { flex:1 1 170px; background:#fff; border:1px solid var(--line);
          border-radius:12px; padding:16px; }
  .kpi-v { font-size:26px; font-weight:700; color:var(--blue); }
  .kpi-l { font-size:12px; color:var(--muted); margin-top:2px; }
  .kpi-s { font-size:11px; color:var(--muted); }
  .card { background:#fff; border:1px solid var(--line); border-radius:12px;
           padding:20px 22px; margin:18px 0; }
  .card.warn { border-color:#fcd34d; background:#fffbeb; }
  h2 { font-size:16px; color:var(--blue); margin:0 0 12px; }
  table { width:100%; border-collapse:collapse; font-size:13px; }
  th, td { text-align:left; padding:7px 9px; border-bottom:1px solid var(--line); }
  th { background:#f1f5f9; color:var(--muted); font-weight:600; }
  td.ess { color:#dc2626; font-weight:600; }
  td.neg { color:#dc2626; } td.pos { color:#16a34a; }
  .grid2 { display:grid; grid-template-columns:1fr 1fr; gap:18px; }
  .verdict { font-weight:700; font-size:15px; }
  .foot { color:var(--muted); font-size:12px; margin-top:24px; }
  .tabs { display:flex; gap:6px; background:#e2e8f0; padding:6px; border-radius:12px;
          margin:22px 0 4px; flex-wrap:wrap; }
  .tabbtn { flex:1 1 180px; border:0; background:transparent; padding:12px 14px;
            border-radius:9px; font-size:14px; font-weight:600; color:var(--muted);
            cursor:pointer; }
  .tabbtn.active { background:#fff; color:var(--blue); box-shadow:0 1px 3px rgba(0,0,0,.08); }
  .tabbtn small { display:block; font-weight:400; font-size:11px; margin-top:2px; }
  .panel { display:none; }
  .panel.active { display:block; }
  .intro { background:#eff6ff; border:1px solid #bfdbfe; border-radius:12px;
           padding:16px 20px; margin:18px 0; font-size:14px; }
  @media (max-width:760px) { .grid2 { grid-template-columns:1fr; } }
</style>"""


def _body_html(res):
    """Inner <main> content for a single scenario (KPIs, charts, tables)."""
    charts = figures.build_all(res)
    cal = _cal_unit(res.unit)
    esc = html.escape

    def img(key, alt):
        return (f'<img alt="{esc(alt)}" '
                f'src="data:image/png;base64,{charts[key]}" style="max-width:100%">')

    claim_html = ""
    if res.claim:
        c = res.claim
        color = "#dc2626" if ("REJECT" in c["verdict"] or "IMPLAUSIBLE" in c["verdict"]) else "#16a34a"
        claim_html = f"""
        <div class="card">
          <h2>Claim check</h2>
          <p>A claim of <b>{c['claim_calendar']:.2f} {cal}</b> implies a
             <b>{c['implied_speedup']:.2f}x</b> whole-project speedup.</p>
          <p>Generous Khan Ceiling: <b>{_fmt_ceiling(res.ceiling_generous)}x</b>.
             Probability the project lands at or below the claim (with scope risk):
             <b>{c['prob_at_or_below_total']:.4f}</b>.</p>
          <p class="verdict" style="color:{color}">Verdict: {esc(c['verdict'])}</p>
        </div>"""

    warn_html = ""
    if res.warnings:
        items = "".join(f"<li>{esc(w)}</li>" for w in res.warnings)
        warn_html = f'<div class="card warn"><h2>Parser notes</h2><ul>{items}</ul></div>'

    # KPI cards
    def kpi(label, value, sub=""):
        return (f'<div class="kpi"><div class="kpi-v">{value}</div>'
                f'<div class="kpi-l">{esc(label)}</div>'
                f'{f"<div class=kpi-s>{esc(sub)}</div>" if sub else ""}</div>')

    kpis = "".join([
        kpi("AI-adjusted calendar", f"{res.adjusted_calendar:.2f}", f"{cal} (was {res.baseline_calendar:.2f})"),
        kpi("Net reduction", f"{res.net_reduction_pct:.1f}%", f"{res.speedup:.2f}x speedup"),
        kpi("Khan Ceiling", f"{_fmt_ceiling(res.ceiling_strict)}-{_fmt_ceiling(res.ceiling_generous)}x",
            f"phi = {res.phi:.2f}"),
        kpi("Absolute floor", f"{res.floor_calendar:.2f}", f"{cal} - nothing beats this"),
        kpi("Total P10-P90", f"{res.total_p10_cal:.1f}-{res.total_p90_cal:.1f}", f"{cal} with scope risk"),
    ])

    # subsystem table
    sub_rows = "".join(
        f"<tr><td>{esc(s['subsystem'])}</td><td>{s['items']}</td>"
        f"<td>{s['baseline']:.2f}</td><td>{s['adjusted']:.2f}</td>"
        f"<td>{s['speedup']:.2f}x</td><td>{s['net_reduction_pct']:.1f}%</td>"
        f"<td>{s['essential_pct']:.0f}%</td></tr>"
        for s in res.subsystems)

    # item table
    item_rows = "".join(
        f"<tr><td>{esc(it.subsystem)}</td><td>{esc(it.name)}</td>"
        f"<td>{esc(it.work_class)}</td><td>{esc(it.difficulty or '-')}</td>"
        f"<td class='{'ess' if it.essential else ''}'>{'Yes' if it.essential else 'No'}</td>"
        f"<td>{it.baseline:.2f}</td><td>{it.a_mode:.2f}</td>"
        f"<td>{it.baseline*it.a_mode:.2f}</td></tr>"
        for it in res.items)

    # P2 table
    p2_rows = "".join(
        f"<tr><td>{esc(m['item'])}</td><td>{m['a_mode']:.2f}</td>"
        f"<td>{m['flat_effort']:.2f}</td><td>{m['true_effort']:.2f}</td>"
        f"<td class='{'neg' if m['misallocation']>0 else 'pos'}'>{m['misallocation']:+.2f}</td></tr>"
        for m in res.misallocation_rows)

    return f"""
  <div class="kpis">{kpis}</div>

  <div class="grid2">
    <div class="card">{img('summary','summary')}</div>
    <div class="card">{img('ceiling','ceiling')}</div>
  </div>

  <div class="card">
    <h2>How the estimate was reached</h2>
    <p>Baseline effort of <b>{res.baseline_effort:.2f} {esc(res.unit)}</b> was partitioned
       into <b>{len(res.items)}</b> work items across <b>{len(res.subsystems)}</b> subsystems.
       Each item carries an AI effort multiplier a<sub>i</sub> (a&lt;1 faster, &gt;1 slower).
       The AI-adjusted effort ratio is r = &Sigma; f<sub>i</sub>a<sub>i</sub> =
       <b>{res.r:.3f}</b>, giving a whole-project speedup S = 1/r =
       <b>{res.speedup:.2f}x</b> and a net effort reduction of
       <b>{res.net_reduction_pct:.1f}%</b>.</p>
    <p>With <b>{res.team_size:g}</b> developer(s), calendar time moves from
       <b>{res.baseline_calendar:.2f}</b> to <b>{res.adjusted_calendar:.2f} {cal}</b>
       (saving {res.calendar_saved:.2f} {cal}). The essential (non-accelerable) fraction is
       &phi; = <b>{res.phi:.2f}</b>, so no amount of AI can push the project below
       <b>{res.floor_calendar:.2f} {cal}</b> (the Khan Ceiling, S<sub>max</sub> =
       {_fmt_ceiling(res.ceiling_strict)}-{_fmt_ceiling(res.ceiling_generous)}x).</p>
  </div>

  {claim_html}

  <div class="card">
    <h2>Uncertainty</h2>
    {img('uncertainty','uncertainty')}
    <p>Conditional band (fixed scope): <b>{res.cond_p10_cal:.2f}-{res.cond_p90_cal:.2f} {cal}</b>.
       Total band including scope/baseline risk:
       <b>{res.total_p10_cal:.2f}-{res.total_p90_cal:.2f} {cal}</b>.</p>
  </div>

  <div class="card">
    <h2>Per-subsystem breakdown</h2>
    {img('subsystems','subsystems')}
    <table><thead><tr><th>Subsystem</th><th>Items</th><th>Baseline</th>
      <th>Adjusted</th><th>Speedup</th><th>Net red.</th><th>Essential</th></tr></thead>
      <tbody>{sub_rows}</tbody></table>
  </div>

  <div class="card">
    <h2>Flat-discount error (P2)</h2>
    <p>A single blanket discount would misallocate <b>{res.misallocated_effort:.2f}
       {esc(res.unit)}</b> across classes even if its total happened to be right.
       Largest offenders:</p>
    <table><thead><tr><th>Work item</th><th>a_mode</th><th>Flat budget</th>
      <th>True need</th><th>Misallocation</th></tr></thead>
      <tbody>{p2_rows}</tbody></table>
  </div>

  <div class="card">
    <h2>All work items</h2>
    <table><thead><tr><th>Subsystem</th><th>Work item</th><th>Class</th>
      <th>Difficulty</th><th>Essential</th><th>Baseline</th><th>a</th><th>Adjusted</th></tr></thead>
      <tbody>{item_rows}</tbody></table>
  </div>

  {warn_html}

  <p class="foot">Generated by the WCE Estimator. Method: Work-Class Estimation with
     the Khan Ceiling (Viquar Khan). Multiplier priors anchored to published field
     studies; uncertainty from 200,000 Monte Carlo draws (seed 20260904). This is a
     decision-support estimate, not a guarantee.</p>
"""


def _page(title, subtitle, inner):
    """Wrap inner <main> content in the full HTML page shell."""
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{html.escape(title)}</title>{_STYLE}</head><body>'
            f'<header><h1>{html.escape(title)}</h1><p>{subtitle}</p></header>'
            f'<main>{inner}</main></body></html>')


_SUBTITLE = ("Work-Class Estimation (WCE) with the Khan Ceiling &middot; "
             "AI-adjusted software effort")


def build_html_report(res, project_name=None):
    """Single-scenario self-contained HTML report."""
    title = project_name or "WCE Estimation Report"
    return _page(title, _SUBTITLE, _body_html(res))


def _tabbed_page(title, tabs, intro=None):
    """
    Assemble a tabbed page from tabs = [(label, sublabel, inner_html), ...].
    First tab is active by default.
    """
    esc = html.escape
    btns, panels = [], []
    for i, (label, sub, inner) in enumerate(tabs):
        active = " active" if i == 0 else ""
        btns.append(
            f'<button class="tabbtn{active}" data-tab="{i}" onclick="wceTab({i})">'
            f'{esc(label)}<small>{esc(sub)}</small></button>')
        panels.append(f'<div class="panel{active}" id="wce-panel-{i}">{inner}</div>')
    intro_html = f'<div class="intro">{intro}</div>' if intro else ""
    inner = (f'{intro_html}<div class="tabs">{"".join(btns)}</div>'
             f'{"".join(panels)}'
             f'<script>function wceTab(n){{'
             f'document.querySelectorAll(".tabbtn").forEach((b,i)=>'
             f'b.classList.toggle("active",i===n));'
             f'document.querySelectorAll(".panel").forEach((p,i)=>'
             f'p.classList.toggle("active",i===n));'
             f'window.scrollTo(0,0);}}</script>')
    return _page(title, _SUBTITLE, inner)


def build_html_report_multi(scenarios, project_name=None, intro=None):
    """
    One self-contained HTML page with a tab per scenario.

    scenarios : list of (label, sublabel, EstimationResult) tuples. The first
                scenario is shown by default.
    intro     : optional HTML string shown above the tabs (e.g. how to read them).
    """
    title = project_name or "WCE Estimation Report"
    tabs = [(label, sub, _body_html(res)) for label, sub, res in scenarios]
    return _tabbed_page(title, tabs, intro=intro)


def _to_months(cal_value, unit, dpm):
    """Convert a calendar figure (in the workbook's unit) to calendar months,
    or None if the unit cannot be converted (e.g. story points)."""
    u = unit.lower()
    if "month" in u:
        return cal_value
    if "day" in u:
        return cal_value / dpm if dpm else None
    if "hour" in u:
        return cal_value / (dpm * 7.0) if dpm else None
    return None


def _months_section(scenarios, esc):
    """Prominent 'calendar months with N developers' block for the landing page."""
    base_res = scenarios[0][2]
    unit = base_res.unit
    dpm = getattr(base_res, "days_per_month", 21.0)
    team = base_res.team_size
    bm = _to_months(base_res.baseline_calendar, unit, dpm)
    if bm is None:
        return ""  # unit not convertible to months; skip the block

    already_months = "month" in unit.lower()
    conv_note = ("" if already_months else
                 f" (assuming ~{dpm:g} working days per month)")

    def kpi(label, value, sub, accent):
        return (f'<div class="kpi"><div class="kpi-v" style="color:{accent}">{value}</div>'
                f'<div class="kpi-l">{esc(label)}</div><div class="kpi-s">{esc(sub)}</div></div>')

    cards = [kpi("Start", f"{bm:.1f} mo", f"{team:g} devs, before AI", "#64748b")]
    accents = ["#16a34a", "#2563eb", "#d97706", "#dc2626"]
    lines = []
    for i, (label, _s, res) in enumerate(scenarios):
        am = _to_months(res.adjusted_calendar, unit, dpm)
        lo = _to_months(res.total_p10_cal, unit, dpm)
        hi = _to_months(res.total_p90_cal, unit, dpm)
        cards.append(kpi(label, f"{am:.1f} mo",
                         f"-{res.net_reduction_pct:.0f}% (range {lo:.1f}-{hi:.1f})",
                         accents[i % len(accents)]))
        lines.append(f"<b>{esc(label)}</b> about <b>{am:.1f} months</b> "
                     f"(P10-P90 {lo:.1f}-{hi:.1f})")
    sentence = (f"With <b>{team:g} developer(s)</b>{conv_note}: the project starts at "
                f"about <b>{bm:.1f} months</b>, then " + ", ".join(lines) + ".")
    return (f'<div class="card" style="border:2px solid #1e3a8a">'
            f'<h2>Bottom line: calendar months with {team:g} developer(s)</h2>'
            f'<div class="kpis" style="margin-top:0">{"".join(cards)}</div>'
            f'<p>{sentence}</p></div>')


def _overview_html(scenarios, cal, esc):
    """
    Landing panel: initial estimate vs each scenario, KPIs + comparison chart +
    a side-by-side table. scenarios = [(label, sublabel, res), ...].
    """
    base_res = scenarios[0][2]
    baseline_cal = base_res.baseline_calendar
    unit = base_res.unit
    chart = figures.fig_compare(baseline_cal, [(l, r) for l, _s, r in scenarios], unit)

    def kpi(label, value, sub, accent="#1e3a8a"):
        return (f'<div class="kpi"><div class="kpi-v" style="color:{accent}">{value}</div>'
                f'<div class="kpi-l">{esc(label)}</div><div class="kpi-s">{esc(sub)}</div></div>')

    cards = [kpi("Initial estimate", f"{baseline_cal:.1f}", f"{cal} (before AI)", "#64748b")]
    accents = ["#16a34a", "#2563eb", "#d97706", "#dc2626"]
    for i, (label, _sub, res) in enumerate(scenarios):
        cards.append(kpi(f"{label}", f"{res.adjusted_calendar:.1f}",
                         f"{cal} (-{res.net_reduction_pct:.0f}%)",
                         accents[i % len(accents)]))
    kpis = "".join(cards)

    # comparison table
    def row(label, res=None):
        if res is None:
            return (f"<tr><td><b>Initial estimate</b></td><td>{baseline_cal:.1f}</td>"
                    f"<td>{base_res.baseline_effort:.1f}</td><td>-</td><td>-</td><td>-</td></tr>")
        return (f"<tr><td><b>{esc(label)}</b></td>"
                f"<td>{res.adjusted_calendar:.1f}</td>"
                f"<td>{res.adjusted_effort:.1f}</td>"
                f"<td>-{res.net_reduction_pct:.0f}%</td>"
                f"<td>{res.speedup:.2f}x</td>"
                f"<td>{res.total_p10_cal:.1f}-{res.total_p90_cal:.1f}</td></tr>")
    trows = row(None, None) + "".join(row(l, r) for l, _s, r in scenarios)

    saved = [f"{l}: {r.calendar_saved:.1f} {cal} saved (-{r.net_reduction_pct:.0f}%)"
             for l, _s, r in scenarios]
    saved_line = " &middot; ".join(esc(x) for x in saved)

    months_block = _months_section(scenarios, esc)

    return f"""
  {months_block}
  <div class="kpis">{kpis}</div>
  <div class="card">
    <h2>Initial vs AI-adjusted</h2>
    <img alt="scenario comparison"
         src="data:image/png;base64,{chart}" style="max-width:100%">
    <p>Starting from an initial estimate of <b>{baseline_cal:.1f} {cal}</b>
       ({base_res.baseline_effort:.1f} {esc(unit)}). {saved_line}. The whiskers show
       the P10-P90 range once scope and estimate uncertainty are included.</p>
  </div>
  <div class="card">
    <h2>Side by side</h2>
    <table><thead><tr><th>Scenario</th><th>Calendar ({cal})</th>
      <th>Effort ({esc(unit)})</th><th>Reduction</th><th>Speedup</th>
      <th>P10-P90 ({cal})</th></tr></thead>
      <tbody>{trows}</tbody></table>
    <p class="foot">Open the <b>{esc(scenarios[0][0])}</b> or <b>{esc(scenarios[-1][0])}</b>
       tab above for the full per-item breakdown, Khan Ceiling, flat-discount error,
       and uncertainty detail behind each number.</p>
  </div>
  {_concepts_html()}"""


def _concepts_html():
    """Plain-language explainer of the WCE method and the Khan Ceiling, with the
    formulas and how each reported number is computed. Shown on the Overview tab."""
    return """
  <div class="card">
    <h2>The concept: Work-Class Estimation (WCE) and the Khan Ceiling</h2>
    <p>AI does not speed up a project by one flat percentage. It speeds up
       different <b>kinds</b> of work by very different amounts, and some work it
       does not speed up at all. WCE splits the estimate into work classes, applies
       a realistic AI multiplier to each, and adds them back up. The
       <b>Khan Ceiling</b> is the hard limit this creates: a project can never go
       faster than its non-accelerable (essential) part allows.</p>

    <h3 style="color:#1e3a8a;font-size:14px;margin:16px 0 6px">The multiplier</h3>
    <p>Every work item gets an AI <b>effort multiplier</b> a: how long it takes
       with AI versus without.</p>
    <table style="max-width:520px">
      <thead><tr><th>a value</th><th>Meaning</th></tr></thead>
      <tbody>
        <tr><td><b>a &lt; 1</b></td><td>AI makes it faster (e.g. a = 0.55 &rarr; 45% faster)</td></tr>
        <tr><td><b>a = 1</b></td><td>no change</td></tr>
        <tr><td><b>a &gt; 1</b></td><td>AI makes it slower (e.g. review/rework, gnarly legacy)</td></tr>
      </tbody>
    </table>

    <h3 style="color:#1e3a8a;font-size:14px;margin:16px 0 6px">The formulas</h3>
    <div style="background:#f8fafc;border:1px solid #e2e8f0;border-radius:10px;
                padding:14px 18px;font-family:ui-monospace,Consolas,monospace;
                font-size:13px;line-height:1.9">
      f&#8341; = share of baseline effort in work class i &nbsp;(&Sigma; f&#8341; = 1)<br>
      a&#8341; = AI effort multiplier for work class i<br>
      <b>r&nbsp; = &Sigma; f&#8341; &middot; a&#8341;</b> &nbsp;&nbsp; &larr; adjusted effort ratio (weighted blend)<br>
      <b>E&#8339;&#8407; = E &middot; r</b> &nbsp;&nbsp; &larr; AI-adjusted effort (E = baseline effort)<br>
      <b>S&nbsp; = 1 / r</b> &nbsp;&nbsp; &larr; whole-project speedup<br>
      <b>S&#8344;&#8342;&#8339; = 1 / &phi;</b> &nbsp;&nbsp; &larr; <b>Khan Ceiling</b> (&phi; = essential, non-accelerable share)<br>
      <b>M&nbsp; = E &middot; &Sigma; f&#8341; &middot; |a&#8341; &minus; r|</b> &nbsp;&nbsp; &larr; flat-discount misallocation (P2)<br>
      <b>G&nbsp; = r / a_focus</b> &nbsp;&nbsp; &larr; codegen-illusion gap (P3)
    </div>

    <h3 style="color:#1e3a8a;font-size:14px;margin:16px 0 6px">How each number here is calculated</h3>
    <ul>
      <li><b>AI-adjusted estimate</b> = baseline &times; r. Each item's effort is
          multiplied by its own a, then summed, so mixed work is priced correctly.</li>
      <li><b>Net reduction</b> = 1 &minus; r. <b>Speedup</b> = 1 / r.</li>
      <li><b>Khan Ceiling (S&#8344;&#8342;&#8339; = 1/&phi;)</b>: the essential work (requirements,
          architecture, coordination, compliance, sign-off) sets a floor. If 25% of
          effort is essential, the project cannot beat 4&times; no matter how good AI is.
          The <b>absolute floor</b> is baseline &times; &phi;.</li>
      <li><b>Uncertainty (P10-P90)</b>: 200,000 Monte Carlo draws vary each
          multiplier over a low/likely/high band and add scope risk, giving a range
          rather than a single false-precision number.</li>
      <li><b>Flat-discount error (P2)</b>: how much effort a single blanket "AI makes
          us X% faster" discount misallocates across classes, even when its total is right.</li>
      <li><b>Codegen illusion (P3)</b>: watching only the fastest task (code
          generation) overstates the whole-project speedup by the factor G.</li>
    </ul>
    <p class="foot">Method and theorem: Work-Class Estimation with the Khan Ceiling,
       Viquar Khan. This tool reproduces the paper's model exactly (same multiplier
       priors, same 200,000-draw Monte Carlo, seed 20260904).</p>
  </div>"""


def build_comparison_report(scenarios, project_name=None, intro=None):
    """
    Landing 'Overview' tab (initial vs scenarios + chart + table) followed by a
    full detail tab per scenario. scenarios = [(label, sublabel, res), ...].
    """
    title = project_name or "WCE Estimation Report"
    esc = html.escape
    cal = _cal_unit(scenarios[0][2].unit)
    overview = _overview_html(scenarios, cal, esc)
    tabs = [("Overview", "initial vs realistic vs optimistic", overview)]
    tabs += [(label, sub, _body_html(res)) for label, sub, res in scenarios]
    return _tabbed_page(title, tabs, intro=intro)
